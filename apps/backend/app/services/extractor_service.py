import io
from typing import Any


class DocumentValidationError(Exception):
    """Base exception for document validation errors."""


class InvalidFormatError(DocumentValidationError):
    """Raised when file content does not match expected format magic bytes."""


class CorruptDocumentError(DocumentValidationError):
    """Raised when document file is malformed or corrupted."""


class EmptyExtractedTextError(DocumentValidationError):
    """Raised when document contains no extractable text."""


def extract_text_from_pdf(content: bytes) -> str:
    from pypdf import PdfReader

    if not content or len(content) == 0:
        raise InvalidFormatError("File PDF kosong")

    if not content.startswith(b"%PDF-"):
        raise InvalidFormatError("File bukan merupakan dokumen PDF yang valid")

    try:
        reader = PdfReader(io.BytesIO(content))
        pages_text: list[str] = []
        for i, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if text.strip():
                pages_text.append(text.strip())
        
        extracted = "\n\n".join(pages_text).strip()
        if not extracted:
            raise EmptyExtractedTextError("Dokumen PDF tidak memuat teks yang dapat diekstrak atau dinilai")
        return extracted
    except (InvalidFormatError, EmptyExtractedTextError):
        raise
    except Exception as e:
        raise CorruptDocumentError(f"File PDF rusak atau tidak dapat dibaca: {e!s}")


def extract_pages_from_pdf(content: bytes) -> list[dict[str, Any]]:
    from pypdf import PdfReader

    if not content or len(content) == 0:
        raise InvalidFormatError("File PDF kosong")

    if not content.startswith(b"%PDF-"):
        raise InvalidFormatError("File bukan merupakan dokumen PDF yang valid")

    try:
        reader = PdfReader(io.BytesIO(content))
        pages: list[dict[str, Any]] = []
        for i, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            pages.append({
                "page": i + 1,
                "text": text.strip(),
            })
        return pages
    except (InvalidFormatError, EmptyExtractedTextError):
        raise
    except Exception as e:
        raise CorruptDocumentError(f"File PDF rusak atau tidak dapat dibaca: {e!s}")


def extract_text_from_docx(content: bytes) -> str:
    import docx

    if not content or len(content) == 0:
        raise InvalidFormatError("File DOCX kosong")

    # DOCX is a zip package starting with PK\x03\x04
    if not content.startswith(b"PK\x03\x04"):
        raise InvalidFormatError("File bukan merupakan dokumen DOCX yang valid")

    try:
        doc = docx.Document(io.BytesIO(content))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    paragraphs.append(" | ".join(row_texts))
        
        extracted = "\n\n".join(paragraphs).strip()
        if not extracted:
            raise EmptyExtractedTextError("Dokumen DOCX tidak memuat teks yang dapat diekstrak atau dinilai")
        return extracted
    except (InvalidFormatError, EmptyExtractedTextError):
        raise
    except Exception as e:
        raise CorruptDocumentError(f"File DOCX rusak atau tidak dapat dibaca: {e!s}")


def extract_text_from_pptx(content: bytes) -> str:
    if not content or len(content) == 0:
        raise InvalidFormatError("File PPTX kosong")
    if not content.startswith(b"PK\x03\x04"):
        raise InvalidFormatError("File bukan merupakan dokumen PPTX yang valid")

    try:
        import pptx
        prs = pptx.Presentation(io.BytesIO(content))
        slides_text: list[str] = []
        for i, slide in enumerate(prs.slides):
            slide_texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        if paragraph.text.strip():
                            slide_texts.append(paragraph.text.strip())
            if slide_texts:
                slides_text.append(f"[Slide {i + 1}]\n" + "\n".join(slide_texts))
        extracted = "\n\n".join(slides_text).strip()
        if not extracted:
            raise EmptyExtractedTextError("Dokumen PPTX tidak memuat teks yang dapat diekstrak atau dinilai")
        return extracted
    except (InvalidFormatError, EmptyExtractedTextError):
        raise
    except Exception as e:
        raise CorruptDocumentError(f"File PPTX rusak atau tidak dapat dibaca: {e!s}")


def extract_text_from_buffer(content: bytes, file_type: str) -> str:
    kind = file_type.upper().lstrip(".")
    if kind == "PDF":
        return extract_text_from_pdf(content)
    elif kind in ("DOCX", "DOC"):
        return extract_text_from_docx(content)
    elif kind in ("PPT", "PPTX"):
        return extract_text_from_pptx(content)
    elif kind in ("TEXT", "TXT"):
        if not content:
            raise EmptyExtractedTextError("Konten teks kosong")
        try:
            text = content.decode("utf-8").strip()
        except UnicodeDecodeError:
            text = content.decode("utf-8", errors="replace").strip()
        if not text:
            raise EmptyExtractedTextError("Konten teks kosong")
        return text
    else:
        # Fallback decode
        try:
            text = content.decode("utf-8").strip()
            if text:
                return text
        except UnicodeDecodeError:
            pass
        return extract_text_from_pdf(content)

