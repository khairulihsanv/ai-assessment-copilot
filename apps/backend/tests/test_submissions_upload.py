import io
from datetime import datetime, timedelta
import docx
import pypdf
import pytest
from fastapi.testclient import TestClient

from app.main import app
from app.models.enums import AssignmentStatus, ClassRole, Role, SubmissionType
from app.models.models import Assignment, Class, Enrollment, Submission, SubmissionVersion, User
from app.security.passwords import hash_password
from app.security.tokens import create_access_token
from app.services.storage_service import storage


def create_valid_pdf_bytes(text: str = "Ini adalah jawaban tugas mahasiswa dalam format PDF") -> bytes:
    stream_content = f"BT /F1 12 Tf 50 700 Td ({text}) Tj ET"
    stream_len = len(stream_content)
    pdf = f"""%PDF-1.4
1 0 obj
<< /Type /Catalog /Pages 2 0 R >>
endobj
2 0 obj
<< /Type /Pages /Kids [3 0 R] /Count 1 >>
endobj
3 0 obj
<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>
endobj
4 0 obj
<< /Length {stream_len} >>
stream
{stream_content}
endstream
endobj
5 0 obj
<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>
endobj
xref
0 6
0000000000 65535 f 
0000000009 00000 n 
0000000058 00000 n 
0000000115 00000 n 
0000000244 00000 n 
0000000330 00000 n 
trailer
<< /Size 6 /Root 1 0 R >>
startxref
417
%%EOF"""
    return pdf.encode("latin1")


def create_blank_pdf_bytes() -> bytes:
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=100, height=100)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def create_valid_docx_bytes(text: str = "Ini adalah jawaban tugas mahasiswa dalam format DOCX") -> bytes:
    doc = docx.Document()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def create_blank_docx_bytes() -> bytes:
    doc = docx.Document()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
def client():
    return TestClient(app)


@pytest.fixture
def setup_upload_env(db_session):
    import uuid
    uid = uuid.uuid4().hex[:8]

    # Dosen
    dosen = User(
        email=f"dosen_upload_{uid}@dexa.ac.id",
        name=f"Dosen Upload {uid}",
        role=Role.DOSEN,
        passwordHash=hash_password("DosenSecurePass!2026"),
    )
    # Mahasiswa enrolled
    mhs = User(
        email=f"mhs_upload_{uid}@dexa.ac.id",
        name=f"Mahasiswa Upload {uid}",
        role=Role.MAHASISWA,
        passwordHash=hash_password("MhsSecurePass!2026"),
    )
    # Mahasiswa not enrolled
    mhs_other = User(
        email=f"mhs_other_{uid}@dexa.ac.id",
        name=f"Mahasiswa Other {uid}",
        role=Role.MAHASISWA,
        passwordHash=hash_password("MhsSecurePass!2026"),
    )
    db_session.add_all([dosen, mhs, mhs_other])
    db_session.commit()

    cls = Class(
        name=f"Kelas Pemrograman Web {uid}",
        description="Kelas untuk pengujian upload",
        subject="Informatika",
        dosenId=dosen.id,
        enrollmentKey=f"UP{uid[:4].upper()}",
    )
    db_session.add(cls)
    db_session.commit()

    enrollment = Enrollment(
        classId=cls.id,
        mahasiswaId=mhs.id,
        role=ClassRole.STUDENT,
    )
    db_session.add(enrollment)
    db_session.commit()

    # Assignment 1: ANY format
    asg_any = Assignment(
        classId=cls.id,
        title="Tugas Format Bebas",
        instructions="Kumpulkan PDF, DOCX, atau Teks",
        submissionType=SubmissionType.ANY,
        status=AssignmentStatus.PUBLISHED,
        dueDate=datetime.now() + timedelta(days=7),
        allowLateSubmission=False,
    )
    # Assignment 2: PDF only
    asg_pdf = Assignment(
        classId=cls.id,
        title="Tugas Khusus PDF",
        instructions="Hanya menerima PDF",
        submissionType=SubmissionType.PDF,
        status=AssignmentStatus.PUBLISHED,
        dueDate=datetime.now() + timedelta(days=7),
        allowLateSubmission=False,
    )
    # Assignment 3: DOCX only
    asg_docx = Assignment(
        classId=cls.id,
        title="Tugas Khusus DOCX",
        instructions="Hanya menerima DOCX",
        submissionType=SubmissionType.DOCX,
        status=AssignmentStatus.PUBLISHED,
        dueDate=datetime.now() + timedelta(days=7),
        allowLateSubmission=False,
    )
    # Assignment 4: Past due
    asg_expired = Assignment(
        classId=cls.id,
        title="Tugas Kadaluwarsa",
        instructions="Tenggat sudah lewat",
        submissionType=SubmissionType.ANY,
        status=AssignmentStatus.PUBLISHED,
        dueDate=datetime.now() - timedelta(days=2),
        allowLateSubmission=False,
    )
    db_session.add_all([asg_any, asg_pdf, asg_docx, asg_expired])
    db_session.commit()

    mhs_token = create_access_token({"sub": mhs.id, "email": mhs.email, "role": mhs.role.value})
    other_token = create_access_token({"sub": mhs_other.id, "email": mhs_other.email, "role": mhs_other.role.value})
    dosen_token = create_access_token({"sub": dosen.id, "email": dosen.email, "role": dosen.role.value})

    return {
        "dosen": dosen,
        "mhs": mhs,
        "mhs_other": mhs_other,
        "class": cls,
        "asg_any": asg_any,
        "asg_pdf": asg_pdf,
        "asg_docx": asg_docx,
        "asg_expired": asg_expired,
        "mhs_token": mhs_token,
        "other_token": other_token,
        "dosen_token": dosen_token,
    }


def test_valid_pdf_upload(client, setup_upload_env, db_session):
    env = setup_upload_env
    pdf_bytes = create_valid_pdf_bytes("Jawaban analitis bab 1 struktur data")

    response = client.post(
        f"/classes/{env['class'].id}/assignments/{env['asg_any'].id}/submissions",
        headers={"Authorization": f"Bearer {env['mhs_token']}"},
        files={"file": ("jawaban.pdf", pdf_bytes, "application/pdf")},
    )
    assert response.status_code == 201, response.text
    data = response.json()
    assert data["success"] is True
    assert data["versionNumber"] == 1

    # Verify version and extracted content in database
    sub = db_session.query(Submission).filter(Submission.id == data["submissionId"]).first()
    assert sub is not None
    assert sub.activeVersionId == data["versionId"]

    ver = db_session.query(SubmissionVersion).filter(SubmissionVersion.id == data["versionId"]).first()
    assert ver.type == SubmissionType.PDF
    assert "Jawaban analitis bab 1" in ver.content
    assert ver.fileUrl is not None
    assert storage.get_sync(ver.fileUrl) == pdf_bytes


def test_valid_docx_upload(client, setup_upload_env, db_session):
    env = setup_upload_env
    docx_bytes = create_valid_docx_bytes("Pembahasan arsitektur microservices")

    response = client.post(
        f"/classes/{env['class'].id}/assignments/{env['asg_any'].id}/submissions",
        headers={"Authorization": f"Bearer {env['mhs_token']}"},
        files={"file": ("laporan.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 201, response.text
    data = response.json()
    assert data["success"] is True

    ver = db_session.query(SubmissionVersion).filter(SubmissionVersion.id == data["versionId"]).first()
    assert ver.type == SubmissionType.DOCX
    assert "arsitektur microservices" in ver.content


def test_valid_text_submission(client, setup_upload_env, db_session):
    env = setup_upload_env
    response = client.post(
        f"/classes/{env['class'].id}/assignments/{env['asg_any'].id}/submissions",
        headers={"Authorization": f"Bearer {env['mhs_token']}"},
        json={"content": "Ini jawaban teks langsung untuk tugas pemrograman web"},
    )
    assert response.status_code == 201, response.text
    data = response.json()
    assert data["success"] is True

    ver = db_session.query(SubmissionVersion).filter(SubmissionVersion.id == data["versionId"]).first()
    assert ver.type == SubmissionType.TEXT
    assert ver.content == "Ini jawaban teks langsung untuk tugas pemrograman web"


def test_empty_file_rejected(client, setup_upload_env):
    env = setup_upload_env
    response = client.post(
        f"/classes/{env['class'].id}/assignments/{env['asg_any'].id}/submissions",
        headers={"Authorization": f"Bearer {env['mhs_token']}"},
        files={"file": ("kosong.pdf", b"", "application/pdf")},
    )
    assert response.status_code == 400
    assert "kosong" in response.text.lower()


def test_corrupt_file_rejected(client, setup_upload_env):
    env = setup_upload_env
    # Corrupt PDF bytes
    corrupt_pdf = b"%PDF-1.4\nthis is definitely a broken and corrupted pdf stream"
    response = client.post(
        f"/classes/{env['class'].id}/assignments/{env['asg_any'].id}/submissions",
        headers={"Authorization": f"Bearer {env['mhs_token']}"},
        files={"file": ("rusak.pdf", corrupt_pdf, "application/pdf")},
    )
    assert response.status_code == 400
    assert "rusak" in response.text.lower() or "tidak dapat dibaca" in response.text.lower()


def test_blank_document_without_text_rejected(client, setup_upload_env):
    env = setup_upload_env
    blank_pdf = create_blank_pdf_bytes()
    response = client.post(
        f"/classes/{env['class'].id}/assignments/{env['asg_any'].id}/submissions",
        headers={"Authorization": f"Bearer {env['mhs_token']}"},
        files={"file": ("blank.pdf", blank_pdf, "application/pdf")},
    )
    assert response.status_code == 400
    assert "tidak memuat teks" in response.text.lower()


def test_file_too_large_rejected(client, setup_upload_env):
    env = setup_upload_env
    large_file = b"%PDF-1.4\n" + (b"A" * (11 * 1024 * 1024))
    response = client.post(
        f"/classes/{env['class'].id}/assignments/{env['asg_any'].id}/submissions",
        headers={"Authorization": f"Bearer {env['mhs_token']}"},
        files={"file": ("large.pdf", large_file, "application/pdf")},
    )
    assert response.status_code == 400
    assert "10mb" in response.text.lower()


def test_format_mismatch_rejected(client, setup_upload_env):
    env = setup_upload_env
    # Assignment requires PDF only, but DOCX is uploaded
    docx_bytes = create_valid_docx_bytes("Konten docx")
    response = client.post(
        f"/classes/{env['class'].id}/assignments/{env['asg_pdf'].id}/submissions",
        headers={"Authorization": f"Bearer {env['mhs_token']}"},
        files={"file": ("dok.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 400
    assert "hanya menerima format" in response.text.lower() or "pdf" in response.text.lower()

    # Assignment requires PDF only, but pure text is sent
    response_text = client.post(
        f"/classes/{env['class'].id}/assignments/{env['asg_pdf'].id}/submissions",
        headers={"Authorization": f"Bearer {env['mhs_token']}"},
        json={"content": "Hanya teks tanpa file"},
    )
    assert response_text.status_code == 400
    assert "pdf" in response_text.text.lower()


def test_submission_persistence_and_detail_reload(client, setup_upload_env, db_session):
    env = setup_upload_env
    pdf_bytes = create_valid_pdf_bytes("Versi 1 jawaban tugas")

    # Upload version 1
    res1 = client.post(
        f"/classes/{env['class'].id}/assignments/{env['asg_any'].id}/submissions",
        headers={"Authorization": f"Bearer {env['mhs_token']}"},
        files={"file": ("jawaban_v1.pdf", pdf_bytes, "application/pdf")},
    )
    assert res1.status_code == 201
    sub_id = res1.json()["submissionId"]

    # Upload version 2
    pdf_bytes_v2 = create_valid_pdf_bytes("Versi 2 revisi perbaikan jawaban tugas")
    res2 = client.post(
        f"/classes/{env['class'].id}/assignments/{env['asg_any'].id}/submissions",
        headers={"Authorization": f"Bearer {env['mhs_token']}"},
        files={"file": ("jawaban_v2.pdf", pdf_bytes_v2, "application/pdf")},
    )
    assert res2.status_code == 201
    data2 = res2.json()
    assert data2["versionNumber"] == 2

    # Reload submission detail via GET
    get_res = client.get(
        f"/submissions/{sub_id}",
        headers={"Authorization": f"Bearer {env['mhs_token']}"},
    )
    assert get_res.status_code == 200
    detail = get_res.json()
    assert detail["activeVersion"]["versionNumber"] == 2
    assert "Versi 2 revisi perbaikan" in detail["activeVersion"]["content"]
    assert len(detail["versions"]) == 2


def test_past_due_date_rejected(client, setup_upload_env):
    env = setup_upload_env
    pdf_bytes = create_valid_pdf_bytes("Jawaban terlambat")
    response = client.post(
        f"/classes/{env['class'].id}/assignments/{env['asg_expired'].id}/submissions",
        headers={"Authorization": f"Bearer {env['mhs_token']}"},
        files={"file": ("terlambat.pdf", pdf_bytes, "application/pdf")},
    )
    assert response.status_code == 400
    assert "tenggat waktu" in response.text.lower()


def test_unauthorized_student_cannot_submit(client, setup_upload_env):
    env = setup_upload_env
    pdf_bytes = create_valid_pdf_bytes("Jawaban mahasiswa lain")
    # Not enrolled student
    response = client.post(
        f"/classes/{env['class'].id}/assignments/{env['asg_any'].id}/submissions",
        headers={"Authorization": f"Bearer {env['other_token']}"},
        files={"file": ("mhs_lain.pdf", pdf_bytes, "application/pdf")},
    )
    assert response.status_code == 403
    assert "belum terdaftar" in response.text.lower()
